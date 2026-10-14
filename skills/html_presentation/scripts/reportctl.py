#!/usr/bin/env python3
"""Report Engineering V1.1 deterministic gate controller.
Agent work remains in SKILL.md; this program never impersonates approval or research.
"""
import argparse, contextlib, datetime as dt, hashlib, json, os, pathlib, re, shutil, sys, tempfile
from urllib.parse import urlparse
import yaml
from jsonschema import Draft202012Validator, FormatChecker

HERE=pathlib.Path(__file__).resolve().parent
SCHEMAS=HERE.parent/'schemas'
STAGES=[f'S{i}' for i in range(1,6)]
GATES=[f'G{i}' for i in range(1,6)]
ARTIFACTS={
 'G1':['brief.yaml'],
 'G2':['research-plan.md','inputs/manifest.yaml','evidence.yaml','narrative.md'],
 'G3':['report-plan.yaml','reviews/content-review.yaml'],
 'G4':['design-spec.yaml','reviews/visual-review.md'],
 'G5':['output/index.html','qa/results.json','qa/qa-report.md','reviews/reader-review.yaml'],
}
SCHEMA_MAP={'brief.yaml':'brief','inputs/manifest.yaml':'sources-manifest','evidence.yaml':'evidence','report-plan.yaml':'report-plan','reviews/content-review.yaml':'content-review','reviews/reader-review.yaml':'content-review','design-spec.yaml':'design-spec','qa/results.json':'qa-results','state.json':'state'}
MARKER=re.compile(r'\{\{claim:(C-\d{3,})\}\}')
REPORT_ID=re.compile(r'^[a-z0-9][a-z0-9-]{2,63}$')
UTC=lambda: dt.datetime.now(dt.timezone.utc).isoformat()
class ReportError(Exception):
 def __init__(self,code,message,details=None):self.code=code;self.message=message;self.details=details or []

def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def json_hash(o):return hashlib.sha256(json.dumps(o,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
def safe(root,path):
 if not isinstance(path,str) or not path or path.startswith('/') or re.match(r'^[A-Za-z]:',path) or '..' in pathlib.PurePosixPath(path).parts or '\\' in path:raise ReportError('E_CONTRACT',f'Unsafe relative path: {path}')
 p=(root/path).resolve()
 if not p.is_relative_to(root.resolve()):raise ReportError('E_CONTRACT',f'Path escapes project root: {path}')
 return p

def load(path):
 try:return json.loads(path.read_text('utf-8')) if path.suffix=='.json' else yaml.safe_load(path.read_text('utf-8'))
 except (OSError,UnicodeError,ValueError,yaml.YAMLError) as e:raise ReportError('E_IO',f'Cannot read {path}: {e}')
def atomic(path,data):
 path.parent.mkdir(parents=True,exist_ok=True)
 fd,name=tempfile.mkstemp(prefix='.tmp-',dir=path.parent)
 try:
  with os.fdopen(fd,'w',encoding='utf-8') as fp:
   fp.write(json.dumps(data,ensure_ascii=False,indent=2));fp.flush();os.fsync(fp.fileno())
  os.replace(name,path)
 finally:
  if os.path.exists(name):os.unlink(name)
def save_yaml(path,data):
 path.parent.mkdir(parents=True,exist_ok=True);path.write_text(yaml.safe_dump(data,allow_unicode=True,sort_keys=False),encoding='utf-8')
def schema_check(path,kind):
 data=load(path)
 schema=json.loads((SCHEMAS/f'{kind}.schema.json').read_text('utf-8'))
 errors=sorted(Draft202012Validator(schema,format_checker=FormatChecker()).iter_errors(data),key=lambda e:list(map(str,e.path)))
 return data,[{'path':str(path)+':'+'.'.join(map(str,e.path)),'reason':e.message} for e in errors]
def project(args):
 if not REPORT_ID.fullmatch(args.report_id):raise ReportError('E_INPUT','report id must be lowercase slug of length 3..64')
 root=pathlib.Path(args.root).expanduser().resolve()/args.report_id
 if args.command!='init' and not (root/'state.json').is_file():raise ReportError('E_IO','No state.json; initialize first')
 return root

def read_state(root):
 s,errs=schema_check(root/'state.json','state')
 if errs:raise ReportError('E_CONTRACT','State invalid',errs)
 return s

def stage_for(g):return 'S'+g[1:]
def gate_for(s):return 'G'+s[1:]
def required_paths(root,gate):
 paths=[]
 for i in range(1,int(gate[1:])+1):
  g='G'+str(i)
  for rel in ARTIFACTS[g]:paths.append(rel)
  if i>=3:
   plan=load(root/'report-plan.yaml') if (root/'report-plan.yaml').is_file() else {}
   for item in plan.get('sections',[]):paths.append(item.get('content_path',''))
  if i>=4:
   design=load(root/'design-spec.yaml') if (root/'design-spec.yaml').is_file() else {}
   paths.extend(design.get('preview_paths',[]))
  if i>=2:
   m=load(root/'inputs/manifest.yaml') if (root/'inputs/manifest.yaml').is_file() else {}
   for source in m.get('sources',[]):paths.extend([source.get('raw_path',''),source.get('snapshot_path','')])
 return sorted(set(paths))
def validate_gate(root,gate):
 errs=[];num=int(gate[1:]);data={}
 for rel in required_paths(root,gate):
  try:
   p=safe(root,rel)
   if not p.is_file():errs.append({'path':rel,'reason':'required file missing'});continue
   if rel in SCHEMA_MAP:
    obj,problems=schema_check(p,SCHEMA_MAP[rel]);data[rel]=obj;errs.extend(problems)
  except (ReportError,KeyError,TypeError) as e:errs.append({'path':rel,'reason':str(e)})
 if errs:return errs
 if num>=2:
  manifest=data['inputs/manifest.yaml'];evidence=data['evidence.yaml']
  source_ids=[x['id'] for x in manifest['sources']];ev_source_ids=[x['id'] for x in evidence['sources']]
  if len(source_ids)!=len(set(source_ids)) or len(ev_source_ids)!=len(set(ev_source_ids)) or set(source_ids)!=set(ev_source_ids):errs.append({'path':'evidence.sources','reason':'duplicate or unmatched source IDs'})
  for s in evidence['sources']:
   if s['manifest_ref']!=s['id']:errs.append({'path':s['id'],'reason':'manifest_ref mismatch'})
  for s in manifest['sources']:
   for filekey,hashkey in [('raw_path','original_sha256'),('snapshot_path','snapshot_sha256')]:
    if digest(safe(root,s[filekey]))!=s[hashkey]:errs.append({'path':s[filekey],'reason':'Source hash mismatch'})
  evids={x['id']:x for x in evidence['evidence']}
  claims={x['id']:x for x in evidence['claims']}
  if len(evids)!=len(evidence['evidence']) or len(claims)!=len(evidence['claims']):errs.append({'path':'evidence.yaml','reason':'duplicate evidence or claim ID'})
  for e in evids.values():
   if e['source_ref'] not in source_ids:errs.append({'path':e['id'],'reason':'dangling source reference'})
  for c in claims.values():
   for ref in c['evidence_refs']:
    if ref not in evids:errs.append({'path':c['id'],'reason':f'dangling evidence {ref}'})
   if c['kind']=='factual' and (num>=3 or c['critical']):
    if c['verification']!='verified' or not c.get('verification_result') or c['verification_result'].get('verdict')!='verified':errs.append({'path':c['id'],'reason':'factual claim not verified'})
    elif not c['evidence_refs'] or set(c['verification_result'].get('evidence_refs',[]))!=set(c['evidence_refs']):errs.append({'path':c['id'],'reason':'verification evidence mismatch'})
   if c['kind']=='factual' and c['verification']=='verified':
    for ref in c['evidence_refs']:
     if ref in evids:
      source=next((s for s in manifest['sources'] if s['id']==evids[ref]['source_ref']),None)
      if source is None or source['extraction_status']!='ok':errs.append({'path':c['id'],'reason':'verified factual with incomplete source extraction'})
  if num>=3:
   plan=data['report-plan.yaml'];sections=plan['sections'];ids=[s['id'] for s in sections]
   if len(set(ids))!=len(ids):errs.append({'path':'report-plan.yaml','reason':'duplicate section IDs'})
   paths=[s['content_path'] for s in sections]
   if len(set(paths))!=len(paths):errs.append({'path':'report-plan.yaml','reason':'duplicate content paths'})
   for s in sections:
    text=safe(root,s['content_path']).read_text('utf-8')
    declared=set(s['claim_refs']);actual=set(MARKER.findall(text))
    if actual!=declared:errs.append({'path':s['content_path'],'reason':f'claim markers mismatch: declared={sorted(declared)}, actual={sorted(actual)}'})
    for ref in actual:
     if ref not in claims:errs.append({'path':s['content_path'],'reason':f'unknown claim {ref}'})
     elif claims[ref]['kind']=='factual' and claims[ref]['verification']!='verified':errs.append({'path':s['content_path'],'reason':f'unverified factual {ref}'})
   cov={c for s in sections for c in s['covers']}
   if not {'desired_outcome','desired_action','core_thesis'}.issubset(cov):errs.append({'path':'report-plan.yaml','reason':'required brief goals not covered'})
   review=data['reviews/content-review.yaml']
   if review['review_kind']!='content' or review['open_blockers'] or any(c['status']!='pass' for c in review['checks']):errs.append({'path':'reviews/content-review.yaml','reason':'content review not passed'})
 if num>=4:
  design=data['design-spec.yaml']
  if len(design['preview_paths'])<2:errs.append({'path':'design-spec.yaml','reason':'at least two previews required'})
 if num>=5:
  html=safe(root,'output/index.html').read_text('utf-8')
  if re.search(r'<(?:script|img|iframe|link|source)\b[^>]*(?:src|href)\s*=\s*["\']https?://',html,re.I) or re.search(r'@import|\bfetch\s*\(|<iframe\b|\bXMLHttpRequest\b',html,re.I):errs.append({'path':'output/index.html','reason':'potential automatic network resource or request'})
  if MARKER.search(html):errs.append({'path':'output/index.html','reason':'unrendered claim marker'})
  qa=data['qa/results.json']
  if qa['qa_status']!='passed' or qa['html_sha256']!=digest(safe(root,'output/index.html')):errs.append({'path':'qa/results.json','reason':'QA not passed or HTML SHA mismatch'})
  actual_viewports={(v['width'],v['height']) for v in qa['viewports'] if v['status']=='pass'}
  if not {(1440,900),(1280,800),(390,844)}.issubset(actual_viewports):errs.append({'path':'qa/results.json','reason':'required viewports not verified'})
  if not qa['checks'] or any(c['status']!='pass' for c in qa['checks']):errs.append({'path':'qa/results.json','reason':'QA checks missing or not passed'})
  if any(f['severity'] in ('P0','P1') and f['status']!='resolved' for f in qa['findings']):errs.append({'path':'qa/results.json','reason':'unresolved P0/P1 finding'})
  reader=data['reviews/reader-review.yaml']
  if reader['review_kind']!='reader' or reader['open_blockers'] or any(c['status']!='pass' for c in reader['checks']):errs.append({'path':'reviews/reader-review.yaml','reason':'reader review not passed'})
 return errs

def snapshot(root,gate,approved):
 files=[]
 for rel in required_paths(root,gate):files.append({'path':rel,'sha256':digest(safe(root,rel))})
 return {'schema_version':'1.1','gate':gate,'artifacts':files,'upstream_approvals':approved}
def sync(root,state):
 changed=None
 for i,g in enumerate(GATES):
  stage=state['stages'][stage_for(g)]
  if stage['status'] not in ('approved','awaiting_approval'):continue
  sub=root/'reviews'/f'{g}-submission.json'
  try:
   old=load(sub);now=snapshot(root,g,{p:state['stages'][stage_for(p)]['approved_digest'] for p in GATES[:i]})
   if old.get('submitted_digest')!=json_hash(now):changed=i;break
  except (ReportError,OSError,KeyError,TypeError):changed=i;break
 if changed is None:return False
 for g in GATES[changed:]:
  st=state['stages'][stage_for(g)]
  if st['status']!='not_started':st['status']='stale';st['submitted_digest']=None;st['approved_digest']=None
 state['release_status']='stale' if state['release_status']=='released' else state['release_status']
 return True

def dispatch(args):
 root=project(args);cmd=args.command
 if cmd=='init':
  if root.exists():raise ReportError('E_EXISTS','Report already exists')
  root.mkdir(parents=True)
  state={'schema_version':'1.1','report_id':args.report_id,'active_stage':'S1','stages':{s:{'status':'in_progress' if s=='S1' else 'not_started','gate':gate_for(s),'submitted_digest':None,'approved_digest':None,'last_error':None} for s in STAGES},'approvals':[],'release_status':'not_released','state_revision':0}
  atomic(root/'state.json',state);return {'report_path':str(root)}
 state=read_state(root)
 if cmd=='status':return {'state':state}
 if cmd=='validate':
  errors=validate_gate(root,args.gate)
  if errors:raise ReportError('E_CONTRACT','Contract validation failed',errors)
  return {'validated_gate':args.gate}
 if cmd=='sync':
  if sync(root,state):state['state_revision']+=1;atomic(root/'state.json',state)
  return {'state':state}
 # sync before mutation; persistent staleness must not be lost on rejected actions
 if sync(root,state):state['state_revision']+=1;atomic(root/'state.json',state)
 if cmd=='resume':
  for s in STAGES:
   entry=state['stages'][s]
   if entry['status'] in ('stale','failed'):
    if s!='S1' and state['stages'][STAGES[STAGES.index(s)-1]]['status']!='approved':raise ReportError('E_STALE','Restore earlier gate first')
    entry.update(status='in_progress',submitted_digest=None,approved_digest=None,last_error=None);state['active_stage']=s;break
  else:return {'state':state,'next_operation':'none'}
 elif cmd=='start-stage':
  s=args.stage;idx=STAGES.index(s)
  if state['stages'][s]['status']!='not_started':raise ReportError('E_STATE',f'{s} not not_started')
  if idx and state['stages'][STAGES[idx-1]]['status']!='approved':raise ReportError('E_GATE','Previous gate not approved')
  state['stages'][s]['status']='in_progress';state['active_stage']=s
 elif cmd=='submit':
  s=stage_for(args.gate);entry=state['stages'][s]
  if entry['status']!='in_progress':raise ReportError('E_STATE','Stage must be in_progress')
  errors=validate_gate(root,args.gate)
  if errors:raise ReportError('E_CONTRACT','Cannot submit invalid gate',errors)
  index=GATES.index(args.gate)
  up={g:state['stages'][stage_for(g)]['approved_digest'] for g in GATES[:index]}
  packet=snapshot(root,args.gate,up);h=json_hash(packet)
  atomic(root/'reviews'/f'{args.gate}-submission.json',{**packet,'submitted_digest':h})
  entry['submitted_digest']=h;entry['status']='awaiting_approval'
 elif cmd=='approve':
  s=stage_for(args.gate);entry=state['stages'][s]
  if entry['status']!='awaiting_approval':raise ReportError('E_STATE','Gate not awaiting approval')
  if entry['submitted_digest']!=args.digest:raise ReportError('E_HASH','Digest mismatch')
  i=GATES.index(args.gate)
  current=snapshot(root,args.gate,{g:state['stages'][stage_for(g)]['approved_digest'] for g in GATES[:i]})
  if json_hash(current)!=args.digest:raise ReportError('E_HASH','Artifact changed after submit')
  confirm=pathlib.Path(args.confirmation_file).read_text('utf-8')
  if args.gate not in confirm or args.digest not in confirm or not re.search(r'批准|同意通过|approve\b',confirm,re.I):raise ReportError('E_CONFIRM','No explicit approval tied to gate and digest')
  entry['status']='approved';entry['approved_digest']=args.digest
  state['approvals'].append({'gate':args.gate,'actor':'owner','decision':'approved','submitted_digest':args.digest,'confirmation_excerpt':confirm[:500],'timestamp':UTC()})
 elif cmd=='reject':
  s=stage_for(args.gate);entry=state['stages'][s]
  if entry['status']!='awaiting_approval':raise ReportError('E_STATE','Gate not awaiting approval')
  text=pathlib.Path(args.confirmation_file).read_text('utf-8');feedback=pathlib.Path(args.feedback_file).read_text('utf-8')
  if not feedback.strip() or args.gate not in text or entry['submitted_digest'] not in text or not re.search(r'拒绝|驳回|reject\b',text,re.I):raise ReportError('E_CONFIRM','Explicit rejection and feedback required')
  entry['status']='rejected';entry['last_error']=feedback[:500]
  state['approvals'].append({'gate':args.gate,'actor':'owner','decision':'rejected','submitted_digest':entry['submitted_digest'],'confirmation_excerpt':text[:500],'timestamp':UTC()})
 elif cmd=='revise':
  e=state['stages'][args.stage]
  if e['status'] not in ('rejected','stale'):raise ReportError('E_STATE','revise requires rejected/stale')
  idx=STAGES.index(args.stage)
  if idx and state['stages'][STAGES[idx-1]]['status']!='approved':raise ReportError('E_GATE','Previous gate invalid')
  e.update(status='in_progress',submitted_digest=None,approved_digest=None,last_error=None);state['active_stage']=args.stage
 elif cmd=='release':
  if state['release_status']=='released':raise ReportError('E_STATE','Already released')
  if any(state['stages'][s]['status']!='approved' for s in STAGES):raise ReportError('E_GATE','G1-G5 approval required')
  errors=validate_gate(root,'G5')
  if errors:raise ReportError('E_QA','Cannot release without passed QA',errors)
  data={'schema_version':'1.1','report_id':args.report_id,'released_at':UTC(),'html_sha256':digest(root/'output/index.html'),'qa_sha256':digest(root/'qa/results.json'),'approvals':{g:state['stages'][stage_for(g)]['approved_digest'] for g in GATES}}
  atomic(root/'output/release-manifest.json',data);state['release_status']='released'
 elif cmd=='intake':
  if state['stages']['S2']['status']!='in_progress':raise ReportError('E_GATE','intake only in S2')
  src=pathlib.Path(args.file or args.snapshot_file)
  if not src.is_file():raise ReportError('E_INPUT','Source missing')
  if args.url:
   u=urlparse(args.url)
   if u.scheme!='https' or not u.netloc:raise ReportError('E_INPUT','HTTPS URL required')
  ext=src.suffix.lower()
  raw=src.read_bytes()
  if args.url or ext in ('.md','.txt'):
   try:normalized=raw.decode('utf-8-sig')
   except UnicodeError:raise ReportError('E_INPUT','UTF-8 input required')
  elif ext=='.docx':
   try:
    from docx import Document
    from io import BytesIO
    doc=Document(BytesIO(raw));normalized='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
   except Exception as e:raise ReportError('E_INPUT',f'DOCX extraction failed: {e}')
  elif ext=='.pdf':
   try:
    import fitz
    doc=fitz.open(stream=raw,filetype='pdf');normalized='\n'.join(f'[page {n+1}]\n'+p.get_text() for n,p in enumerate(doc))
   except Exception as e:raise ReportError('E_INPUT',f'PDF extraction failed: {e}')
  else:raise ReportError('E_INPUT','Unsupported input format')
  if not normalized.strip():raise ReportError('E_INPUT','Extraction returned no text')
  manifest=root/'inputs/manifest.yaml'
  o=load(manifest) if manifest.exists() else {'schema_version':'1.1','sources':[]}
  ident=f'SRC-{len(o["sources"])+1:03d}'
  while ident in {s['id'] for s in o['sources']}:ident=f'SRC-{int(ident[4:])+1:03d}'
  relraw=f'inputs/raw/{ident}{ext if ext else ".txt"}';relsnap=f'inputs/snapshots/{ident}.txt'
  rawpath=safe(root,relraw);rawpath.parent.mkdir(parents=True,exist_ok=True);rawpath.write_bytes(raw)
  snappath=safe(root,relsnap);snappath.parent.mkdir(parents=True,exist_ok=True);snappath.write_text(normalized,'utf-8')
  o['sources'].append({'id':ident,'origin':args.url or str(src),'kind':'web' if args.url else 'local','raw_path':relraw,'snapshot_path':relsnap,'original_sha256':digest(rawpath),'snapshot_sha256':digest(snappath),'extraction_status':'ok','retrieved_at':UTC(),'classification':args.classification})
  save_yaml(manifest,o)
  response={'source_id':ident,'raw_path':relraw,'snapshot_path':relsnap}
 else:raise ReportError('E_INPUT','Unknown operation')
 state['state_revision']+=1
 atomic(root/'state.json',state)
 return {'state_revision':state['state_revision'],**({'digest':h} if cmd=='submit' else {}),**(response if cmd=='intake' else {}),**({'state':state} if cmd=='resume' else {})}

def main():
 p=argparse.ArgumentParser(description=__doc__);sub=p.add_subparsers(dest='command',required=True)
 for c in ['init','status','validate','start-stage','submit','approve','reject','revise','sync','resume','release','intake']:
  x=sub.add_parser(c);x.add_argument('report_id');x.add_argument('--root',default='projects')
  if c in ('submit','validate','approve','reject'):x.add_argument('--gate',required=True,choices=GATES)
  if c in ('start-stage','revise'):x.add_argument('--stage',required=True,choices=STAGES)
  if c=='approve':x.add_argument('--digest',required=True);x.add_argument('--confirmation-file',required=True)
  if c=='reject':x.add_argument('--feedback-file',required=True);x.add_argument('--confirmation-file',required=True)
  if c=='intake':
   group=x.add_mutually_exclusive_group(required=True);group.add_argument('--file');group.add_argument('--url')
   x.add_argument('--snapshot-file');x.add_argument('--classification',choices=['public','internal'],required=True)
 args=p.parse_args()
 try:
  if args.command=='intake' and bool(args.url)!=bool(args.snapshot_file):raise ReportError('E_INPUT','--url requires --snapshot-file and vice versa')
  result=dispatch(args);print(json.dumps({'ok':True,'code':'OK','message':'success','details':[],**result},ensure_ascii=False));return 0
 except ReportError as e:
  print(json.dumps({'ok':False,'code':e.code,'message':e.message,'details':e.details},ensure_ascii=False))
  return 2 if e.code in ('E_INPUT','E_CONTRACT','E_EXISTS') else 3 if e.code in ('E_STATE','E_GATE','E_CONFIRM','E_HASH','E_STALE') else 4 if e.code in ('E_QA','E_BLOCKED') else 5
 except (OSError,KeyError,TypeError) as e:
  print(json.dumps({'ok':False,'code':'E_IO','message':str(e),'details':[]},ensure_ascii=False));return 5
if __name__=='__main__':sys.exit(main())
